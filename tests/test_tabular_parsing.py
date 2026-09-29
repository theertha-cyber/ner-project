"""Safe file parsing (tabular-file-ingestion spec): cached formula values only,
merged-header rejection, and the row cap enforced while streaming."""

import os
import re
import zipfile

import openpyxl
import pytest

from src.shared.tabular_files import ingest
from src.shared.tabular_files.ingest import TabularIngestError


def _with_cached_formula_value(path: str, cell: str, cached: str) -> None:
    """openpyxl writes formulas without a cached value; Excel writes both. Patch
    the sheet XML so the cell carries `<f>` and a cached `<v>` like a saved
    Excel workbook does."""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                pattern = re.compile(rf'(<c r="{cell}"[^>]*>)<f>([^<]*)</f>(<v\s*/>|<v></v>)?(</c>)')
                xml, count = pattern.subn(rf"\1<f>\2</f><v>{cached}</v>\4", xml)
                assert count == 1, xml
                data = xml.encode("utf-8")
            dst.writestr(item, data)
    os.replace(tmp, path)


def _stage(tmp_path, original, kind, max_rows=1000, sheet=None):
    staging = str(tmp_path / "staging.duckdb")
    labels, used_sheet = ingest.stage_original(str(original), kind, sheet, staging, max_rows, str(tmp_path))
    return staging, labels, used_sheet


def test_formula_cells_use_cached_values(tmp_path, monkeypatch):
    """Scenario: Formula cells use cached values."""
    path = tmp_path / "orders.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Item", "Qty", "Price", "D"])
    ws.append(["Widget", 4, 300, "=B2*C2"])
    wb.save(path)
    _with_cached_formula_value(str(path), "D2", "1200")

    # Evaluating a formula would need a calculation engine; make sure none runs.
    import openpyxl.formula.tokenizer as tokenizer
    monkeypatch.setattr(tokenizer.Tokenizer, "__init__", lambda *a, **k: pytest.fail("formula evaluated"))

    staging, labels, sheet = _stage(tmp_path, path, "xlsx")
    assert labels == ["Item", "Qty", "Price", "D"]
    assert sheet == "Orders"
    con = ingest.open_staging(staging)
    try:
        assert con.execute("SELECT c3 FROM staging").fetchall() == [("1200",)]
    finally:
        con.close()


def test_merged_header_is_rejected(tmp_path):
    """Scenario: Merged header is rejected."""
    path = tmp_path / "budget.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Quarterly budget"
    ws.merge_cells("A1:C2")
    ws.append(["Region", "Amount", "Owner"])
    ws.append(["EMEA", 10, "x"])
    wb.save(path)
    with pytest.raises(TabularIngestError) as exc:
        _stage(tmp_path, path, "xlsx")
    assert exc.value.code == ingest.UNSUPPORTED_SHEET_LAYOUT
    assert not os.path.exists(tmp_path / "staging.duckdb")


def test_missing_header_cell_is_rejected(tmp_path):
    path = tmp_path / "gaps.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Region", None, "Owner"])
    ws.append(["EMEA", 10, "x"])
    wb.save(path)
    with pytest.raises(TabularIngestError) as exc:
        _stage(tmp_path, path, "xlsx")
    assert exc.value.code == ingest.UNSUPPORTED_SHEET_LAYOUT


def test_row_cap_is_enforced_while_streaming_csv(tmp_path, monkeypatch):
    """Scenario: Row cap is enforced while streaming. Scaled: a cap of 1,000 with
    1,200 rows exercises the same `LIMIT cap+1` stop the 1M cap uses."""
    path = tmp_path / "big.csv"
    with open(path, "w") as fh:
        fh.write("id,amount\n")
        for i in range(1200):
            fh.write(f"{i},{i}\n")
    executed = []
    real_connect = ingest._connect

    def spying_connect(p=":memory:"):
        con = real_connect(p)

        class Spy:
            def __getattr__(self, name):
                return getattr(con, name)

            def execute(self, sql, *args):
                executed.append((sql, args))
                return con.execute(sql, *args)

        return Spy()

    monkeypatch.setattr(ingest, "_connect", spying_connect)
    with pytest.raises(TabularIngestError) as exc:
        _stage(tmp_path, path, "csv", max_rows=1000)
    assert exc.value.code == ingest.ROW_LIMIT_EXCEEDED
    # The read itself is bounded: it never asks for more than header + cap + 1 rows.
    read = [args for sql, args in executed if "read_csv" in sql]
    assert read and read[0][0][1] == 1002


def test_row_cap_is_enforced_while_streaming_xlsx(tmp_path):
    path = tmp_path / "big.xlsx"
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    ws.append(["id"])
    for i in range(30):
        ws.append([i])
    wb.save(path)
    with pytest.raises(TabularIngestError) as exc:
        _stage(tmp_path, path, "xlsx", max_rows=20)
    assert exc.value.code == ingest.ROW_LIMIT_EXCEEDED


def test_csv_keeps_original_header_labels(tmp_path):
    path = tmp_path / "sales.csv"
    path.write_text("Order ID,Amount,amount\n1,2,3\n")
    staging, labels, sheet = _stage(tmp_path, path, "csv")
    assert labels == ["Order ID", "Amount", "amount"]
    assert sheet is None
    con = ingest.open_staging(staging)
    try:
        assert con.execute("SELECT __row, c0, c1, c2 FROM staging").fetchall() == [(2, "1", "2", "3")]
    finally:
        con.close()
