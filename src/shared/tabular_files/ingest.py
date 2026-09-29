"""Staging, type inference, profiling, load reports and Parquet publication for
uploaded tabular files (ADR-018 Decision 6, tabular-file-ingestion spec).

Trusted platform code: it runs in the ingest worker (and in the gateway for the
review load report) and may use DuckDB with external access. Nothing here ever
evaluates a formula, a macro or model output.

Staging layout: a worker-local DuckDB file holding one table, `staging`, with a
`__row` column (the record's row number, counting the header as row 1) and one
VARCHAR column per source column, named positionally `c0`, `c1`, ... so that
administrator renames never touch staged data. Header labels are kept apart in
the profile.

Every type decision runs in SQL over the whole staging table — never a sample —
and uses strict regular expressions in front of `TRY_CAST`, because DuckDB's
casts are lenient in ways the product rules are not (`'1.5'::BIGINT` rounds,
`'0x1A'::BIGINT` is 26, a timestamp string casts to a date).
"""

from __future__ import annotations

import csv
import datetime as dt
import decimal
import os
import xml.etree.ElementTree as ET
import zipfile

import duckdb

from src.shared.tabular_files.identifiers import is_valid_identifier, sanitize_all, sanitize_identifier

DEFAULT_NULL_TOKENS = ["", "n/a", "na", "null", "-", "#n/a"]

TYPE_TEXT = "text"
TYPE_BOOLEAN = "boolean"
TYPE_BIGINT = "bigint"
TYPE_NUMERIC = "numeric"
TYPE_DATE = "date"
TYPE_TIMESTAMP = "timestamp"
ALLOWED_TYPES = (TYPE_TEXT, TYPE_BOOLEAN, TYPE_BIGINT, TYPE_NUMERIC, TYPE_DATE, TYPE_TIMESTAMP)
# Inference order: the first candidate every non-null value converts to wins.
CANDIDATE_ORDER = (TYPE_BOOLEAN, TYPE_BIGINT, TYPE_NUMERIC, TYPE_DATE, TYPE_TIMESTAMP)

DATE_FORMAT_ISO = "YYYY-MM-DD"
DATE_FORMAT_DMY = "DD/MM/YYYY"
DATE_FORMAT_MDY = "MM/DD/YYYY"
DATE_FORMATS = (DATE_FORMAT_ISO, DATE_FORMAT_DMY, DATE_FORMAT_MDY)
_STRPTIME = {DATE_FORMAT_DMY: "%d/%m/%Y", DATE_FORMAT_MDY: "%m/%d/%Y"}

# Finite failure reasons, stored on the version row and shown in the portal.
ROW_LIMIT_EXCEEDED = "ROW_LIMIT_EXCEEDED"
UNSUPPORTED_SHEET_LAYOUT = "UNSUPPORTED_SHEET_LAYOUT"
UNPARSEABLE_FILE = "UNPARSEABLE_FILE"
EMPTY_FILE = "EMPTY_FILE"
SHEET_NOT_FOUND = "SHEET_NOT_FOUND"

# Publish gate codes.
DATE_FORMAT_REQUIRED = "DATE_FORMAT_REQUIRED"
DESCRIPTION_REQUIRED = "DESCRIPTION_REQUIRED"
DUPLICATE_IDENTIFIER = "DUPLICATE_IDENTIFIER"
NO_INCLUDED_COLUMNS = "NO_INCLUDED_COLUMNS"

MAX_TOP_VALUES = 10
MAX_DISTINCT_FOR_HINTS = 50
MAX_WARNING_ROWS = 10
MAX_REJECT_ROWS = 100
PREVIEW_ROWS = 20
MAX_NULL_TOKENS = 20
MAX_NULL_TOKEN_LENGTH = 32
MAX_DESCRIPTION_LENGTH = 500

_INT_STRICT = r"[+-]?(0|[1-9][0-9]*)"
_INT_LOOSE = r"[+-]?[0-9]+"
_NUM_STRICT = r"[+-]?(0|[1-9][0-9]*)(\.[0-9]*)?([eE][+-]?[0-9]+)?|[+-]?\.[0-9]+([eE][+-]?[0-9]+)?"
_NUM_LOOSE = r"[+-]?([0-9]+(\.[0-9]*)?|\.[0-9]+)([eE][+-]?[0-9]+)?"
_ISO_DATE = r"[0-9]{4}-[0-9]{2}-[0-9]{2}"
_SLASH_DATE = r"[0-9]{1,2}/[0-9]{1,2}/[0-9]{4}"
_TIMESTAMP = r"[0-9]{4}-[0-9]{2}-[0-9]{2}[ T][0-9]{2}:[0-9]{2}(:[0-9]{2}(\.[0-9]+)?)?"


class TabularIngestError(Exception):
    """A finite, safe failure reason; `str(e)` is the code."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


class ReviewInvalid(Exception):
    """A review edit was refused; names the field, never echoes the value."""

    def __init__(self, field: str, message: str):
        super().__init__(message)
        self.field = field
        self.message = message


# --------------------------------------------------------------------- staging

def _connect(path: str = ":memory:"):
    con = duckdb.connect(path)
    # File order is row order: one thread and preserved insertion order keep
    # `rowid` equal to the source record index.
    con.execute("SET threads=1")
    con.execute("SET preserve_insertion_order=true")
    return con


def _run(con, sql: str, params: dict):
    """DuckDB refuses a named parameter the statement does not use, and a
    statement whose columns are all text never references `$tokens`."""
    return con.execute(sql, {k: v for k, v in params.items() if f"${k}" in sql})


def _col(i: int) -> str:
    return f"c{int(i)}"


def stage_csv(con, csv_path: str, max_rows: int) -> list[str]:
    """Reads a CSV as text-only columns into `staging` and returns the header
    labels. The header record is read as data (`header=false`) so that blank or
    duplicate headers keep their exact original label. `LIMIT` stops the scan
    as soon as the row cap is exceeded, so an oversized file is never loaded
    whole."""
    path = csv_path.replace("\\", "/")
    try:
        con.execute(
            "CREATE OR REPLACE TABLE raw AS SELECT * FROM read_csv(?, all_varchar=true, header=false, "
            "null_padding=true) LIMIT ?",
            [path, max_rows + 2],
        )
    except duckdb.Error as exc:
        raise TabularIngestError(UNPARSEABLE_FILE) from exc
    total = con.execute("SELECT COUNT(*) FROM raw").fetchone()[0]
    if total == 0:
        raise TabularIngestError(EMPTY_FILE)
    if total - 1 > max_rows:
        con.execute("DROP TABLE raw")
        raise TabularIngestError(ROW_LIMIT_EXCEEDED)
    source_columns = [r[0] for r in con.execute("DESCRIBE raw").fetchall()]
    header = con.execute("SELECT * FROM raw WHERE rowid = 0").fetchone()
    labels = ["" if v is None else str(v).strip() for v in header]
    select = ", ".join(f'"{name}" AS {_col(i)}' for i, name in enumerate(source_columns))
    con.execute(
        f"CREATE OR REPLACE TABLE staging AS SELECT CAST(rowid + 1 AS BIGINT) AS __row, {select} "
        "FROM raw WHERE rowid > 0 ORDER BY rowid"
    )
    con.execute("DROP TABLE raw")
    return labels


def _sheet_merged_ranges(xlsx_path: str, worksheet_path: str) -> list[tuple[int, int]]:
    """(min_row, max_row) of every merged range, streamed from the sheet XML —
    read-only worksheets do not expose merged cells."""
    from openpyxl.utils.cell import range_boundaries

    ranges: list[tuple[int, int]] = []
    with zipfile.ZipFile(xlsx_path) as zf, zf.open(worksheet_path) as fh:
        for _, elem in ET.iterparse(fh):
            if elem.tag.endswith("}mergeCell") or elem.tag == "mergeCell":
                ref = elem.get("ref")
                if ref:
                    _, min_row, _, max_row = range_boundaries(ref)
                    ranges.append((min_row, max_row))
            elem.clear()
    return ranges


def _cell_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def xlsx_to_csv(xlsx_path: str, sheet: str | None, out_csv: str, max_rows: int) -> str:
    """Streams one sheet's cached cell values into a text CSV and returns the
    sheet name used. `read_only=True, data_only=True`: formulas are never
    evaluated — a formula cell contributes its cached value only. Rejects the
    sheet when merged cells touch the header region or no header row exists."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as exc:
        raise TabularIngestError(UNPARSEABLE_FILE) from exc
    try:
        if sheet is None:
            ws = wb.worksheets[0]
        elif sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            raise TabularIngestError(SHEET_NOT_FOUND)
        sheet_name = ws.title
        merged = _sheet_merged_ranges(xlsx_path, ws._worksheet_path)

        header_row_number = None
        header: list | None = None
        written = 0
        with open(out_csv, "w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = [_cell_text(v) for v in row]
                if header is None:
                    if not any(v not in (None, "") for v in values):
                        continue
                    header_row_number = row_number
                    # Header region: the header row and anything above it. A
                    # merged title banner sits there, and so does a merged
                    # multi-row header.
                    if any(min_row <= header_row_number for min_row, _ in merged):
                        raise TabularIngestError(UNSUPPORTED_SHEET_LAYOUT)
                    while values and values[-1] in (None, ""):
                        values.pop()
                    if any(v in (None, "") for v in values):
                        raise TabularIngestError(UNSUPPORTED_SHEET_LAYOUT)
                    header = values
                    writer.writerow(header)
                    continue
                if not any(v not in (None, "") for v in values):
                    continue
                written += 1
                if written > max_rows:
                    raise TabularIngestError(ROW_LIMIT_EXCEEDED)
                values = (values + [None] * len(header))[: len(header)]
                writer.writerow(["" if v is None else v for v in values])
        if header is None:
            raise TabularIngestError(UNSUPPORTED_SHEET_LAYOUT)
        return sheet_name
    finally:
        wb.close()


def stage_original(original_path: str, source_kind: str, sheet: str | None, staging_path: str,
                   max_rows: int, scratch_dir: str) -> tuple[list[str], str | None]:
    """Builds the staging DuckDB file from an original upload. Returns the
    header labels and the sheet used (None for CSV)."""
    if os.path.exists(staging_path):
        os.remove(staging_path)
    used_sheet = None
    csv_path = original_path
    if source_kind == "xlsx":
        csv_path = os.path.join(scratch_dir, "sheet.csv")
        used_sheet = xlsx_to_csv(original_path, sheet, csv_path, max_rows)
    con = _connect(staging_path)
    try:
        labels = stage_csv(con, csv_path, max_rows)
        con.execute("CREATE OR REPLACE TABLE staging_labels (idx INTEGER, label VARCHAR)")
        con.executemany("INSERT INTO staging_labels VALUES (?, ?)", list(enumerate(labels)))
    except BaseException:
        con.close()
        if os.path.exists(staging_path):
            os.remove(staging_path)
        raise
    finally:
        if source_kind == "xlsx" and os.path.exists(csv_path):
            os.remove(csv_path)
    con.close()
    return labels, used_sheet


def open_staging(staging_path: str, read_only: bool = True):
    con = duckdb.connect(staging_path, read_only=read_only)
    con.execute("SET threads=1")
    return con


def staged_labels(con) -> list[str]:
    return [r[0] for r in con.execute("SELECT label FROM staging_labels ORDER BY idx").fetchall()]


# ------------------------------------------------------------------ inference

def _nv(col: str) -> str:
    """The null-mapped, trimmed value of one staged column. `$tokens` is the
    lower-cased null-token list, bound as a parameter."""
    return (f"(CASE WHEN {col} IS NULL OR list_contains($tokens, lower(trim({col}))) "
            f"THEN NULL ELSE trim({col}) END)")


def _strict_ok(kind: str, v: str) -> str:
    """Inference predicate: the value is a clean instance of `kind`."""
    if kind == TYPE_BOOLEAN:
        return f"lower({v}) IN ('true', 'false', 'yes', 'no')"
    if kind == TYPE_BIGINT:
        return f"(regexp_full_match({v}, '{_INT_STRICT}') AND TRY_CAST({v} AS BIGINT) IS NOT NULL)"
    if kind == TYPE_NUMERIC:
        return f"(regexp_full_match({v}, '{_NUM_STRICT}') AND TRY_CAST({v} AS DECIMAL(38,9)) IS NOT NULL)"
    if kind == TYPE_DATE:
        return f"(regexp_full_match({v}, '{_ISO_DATE}') AND TRY_CAST({v} AS DATE) IS NOT NULL)"
    if kind == TYPE_TIMESTAMP:
        return (f"((regexp_full_match({v}, '{_TIMESTAMP}') OR regexp_full_match({v}, '{_ISO_DATE}')) "
                f"AND TRY_CAST({v} AS TIMESTAMP) IS NOT NULL)")
    raise ValueError(kind)


def cast_expr(kind: str, date_format: str | None, v: str) -> str:
    """The approved cast of a null-mapped value, NULL when it does not convert.
    Used by the load report and by publish, so the two cannot disagree."""
    if kind == TYPE_TEXT:
        return v
    if kind == TYPE_BOOLEAN:
        return (f"(CASE WHEN lower({v}) IN ('true', 'yes') THEN TRUE "
                f"WHEN lower({v}) IN ('false', 'no') THEN FALSE ELSE NULL END)")
    if kind == TYPE_BIGINT:
        return f"(CASE WHEN regexp_full_match({v}, '{_INT_LOOSE}') THEN TRY_CAST({v} AS BIGINT) END)"
    if kind == TYPE_NUMERIC:
        return f"(CASE WHEN regexp_full_match({v}, '{_NUM_LOOSE}') THEN TRY_CAST({v} AS DECIMAL(38,9)) END)"
    if kind == TYPE_DATE:
        if date_format == DATE_FORMAT_ISO:
            return f"(CASE WHEN regexp_full_match({v}, '{_ISO_DATE}') THEN TRY_CAST({v} AS DATE) END)"
        if date_format in _STRPTIME:
            return (f"(CASE WHEN regexp_full_match({v}, '{_SLASH_DATE}') "
                    f"THEN CAST(try_strptime({v}, '{_STRPTIME[date_format]}') AS DATE) END)")
        raise ValueError("date format required")
    if kind == TYPE_TIMESTAMP:
        return (f"(CASE WHEN regexp_full_match({v}, '{_TIMESTAMP}') OR regexp_full_match({v}, '{_ISO_DATE}') "
                f"THEN TRY_CAST({v} AS TIMESTAMP) END)")
    raise ValueError(kind)


def _duck_type(kind: str) -> str:
    return {TYPE_TEXT: "VARCHAR", TYPE_BOOLEAN: "BOOLEAN", TYPE_BIGINT: "BIGINT",
            TYPE_NUMERIC: "DECIMAL(38,9)", TYPE_DATE: "DATE", TYPE_TIMESTAMP: "TIMESTAMP"}[kind]


def _jsonable(value):
    if isinstance(value, decimal.Decimal):
        normalized = value.normalize()
        return int(normalized) if normalized == normalized.to_integral_value() else float(normalized)
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def infer_column(con, index: int, tokens: list[str]) -> dict:
    """Candidate elimination for one staged column over every row."""
    col = _col(index)
    v = _nv(col)
    params = {"tokens": tokens}
    base = f"SELECT __row, {v} AS v FROM staging"
    non_null, nulls = con.execute(
        f"SELECT COUNT(v), COUNT(*) - COUNT(v) FROM ({base})", params
    ).fetchone()
    result = {"type": TYPE_TEXT, "date_format": None, "date_format_required": False,
              "null_count": int(nulls), "warnings": []}
    if non_null == 0:
        return result

    failures: dict[str, int] = {}
    for kind in CANDIDATE_ORDER:
        failed = con.execute(
            f"SELECT COUNT(*) FROM ({base}) WHERE v IS NOT NULL AND NOT ({_strict_ok(kind, 'v')})",
            params,
        ).fetchone()[0]
        failures[kind] = int(failed)
        if failed == 0:
            result["type"] = kind
            if kind == TYPE_DATE:
                result["date_format"] = DATE_FORMAT_ISO
            return result

    # Slash dates: decide the format from the evidence, never by guessing.
    slash_ok = con.execute(
        f"SELECT COUNT(*) FROM ({base}) WHERE v IS NOT NULL AND regexp_full_match(v, '{_SLASH_DATE}')",
        params,
    ).fetchone()[0]
    if slash_ok == non_null:
        dmy, mdy = con.execute(
            f"SELECT COUNT(try_strptime(v, '%d/%m/%Y')), COUNT(try_strptime(v, '%m/%d/%Y')) "
            f"FROM ({base}) WHERE v IS NOT NULL",
            params,
        ).fetchone()
        if dmy == non_null and mdy == non_null:
            result.update(type=TYPE_DATE, date_format=None, date_format_required=True)
            return result
        if dmy == non_null:
            result.update(type=TYPE_DATE, date_format=DATE_FORMAT_DMY)
            return result
        if mdy == non_null:
            result.update(type=TYPE_DATE, date_format=DATE_FORMAT_MDY)
            return result

    # Falls back to text. When the strictest numeric or temporal candidate failed
    # only on a minority of values, say so, with the first offending rows.
    near = [(kind, n) for kind, n in failures.items() if kind != TYPE_BOOLEAN and 0 < n * 2 < non_null]
    if near:
        kind, count = min(near, key=lambda kn: (kn[1], CANDIDATE_ORDER.index(kn[0])))
        rows = [r[0] for r in con.execute(
            f"SELECT __row FROM ({base}) WHERE v IS NOT NULL AND NOT ({_strict_ok(kind, 'v')}) "
            f"ORDER BY __row LIMIT {MAX_WARNING_ROWS}",
            params,
        ).fetchall()]
        result["warnings"].append({"code": "mixed_values", "candidate": kind, "count": count,
                                   "rows": [int(r) for r in rows]})
    return result


def build_profile(con, labels: list[str], relation: str, null_tokens: list[str] | None = None) -> dict:
    """The draft profile: per-column inference and stats, plus the row count."""
    tokens = [t.lower() for t in (DEFAULT_NULL_TOKENS if null_tokens is None else null_tokens)]
    identifiers = sanitize_all(labels)
    row_count = con.execute("SELECT COUNT(*) FROM staging").fetchone()[0]
    columns = []
    for index, (label, identifier) in enumerate(zip(labels, identifiers)):
        inferred = infer_column(con, index, tokens)
        v = _nv(_col(index))
        params = {"tokens": tokens}
        distinct = con.execute(f"SELECT COUNT(DISTINCT {v}) FROM staging", params).fetchone()[0]
        column = {
            "index": index, "identifier": identifier, "label": label,
            "type": inferred["type"], "date_format": inferred["date_format"],
            "date_format_required": inferred["date_format_required"],
            "null_count": inferred["null_count"], "distinct_count": int(distinct),
            "min": None, "max": None, "top_values": [], "warnings": inferred["warnings"],
        }
        kind = inferred["type"]
        if kind in (TYPE_BIGINT, TYPE_NUMERIC, TYPE_TIMESTAMP) or (
                kind == TYPE_DATE and inferred["date_format"]):
            expr = cast_expr(kind, inferred["date_format"], v)
            lo, hi = con.execute(f"SELECT MIN({expr}), MAX({expr}) FROM staging", params).fetchone()
            column["min"], column["max"] = _jsonable(lo), _jsonable(hi)
        if kind == TYPE_TEXT and 0 < distinct <= MAX_DISTINCT_FOR_HINTS:
            column["top_values"] = [
                {"value": value, "count": int(count)}
                for value, count in con.execute(
                    f"SELECT {v} AS val, COUNT(*) AS n FROM staging WHERE {v} IS NOT NULL "
                    f"GROUP BY val ORDER BY n DESC, val LIMIT {MAX_TOP_VALUES}",
                    params,
                ).fetchall()
            ]
        columns.append(column)
    return {"relation": relation, "row_count": int(row_count), "columns": columns,
            "null_tokens": list(tokens)}


def draft_review(profile: dict) -> dict:
    """The initial review: the profile's choices, nothing excluded, no
    descriptions, every value hint listed but unapproved."""
    return {
        "table": {"relation": profile["relation"], "description": "",
                  "null_tokens": list(profile.get("null_tokens") or DEFAULT_NULL_TOKENS)},
        "columns": [
            {
                "index": c["index"], "label": c["label"], "identifier": c["identifier"],
                "type": c["type"], "date_format": c["date_format"], "excluded": False,
                "description": "",
                "value_hints": [{"value": h["value"], "approved": False} for h in c["top_values"]],
            }
            for c in profile["columns"]
        ],
    }


# -------------------------------------------------------------------- review

def validate_review(edit: dict, profile: dict, current: dict) -> dict:
    """Applies an administrator's review edit on top of the current review and
    returns the new review. Refuses anything malformed with the field name."""
    if not isinstance(edit, dict):
        raise ReviewInvalid("body", "Request body must be a JSON object.")
    unknown = set(edit) - {"table", "columns"}
    if unknown:
        field = sorted(unknown)[0]
        raise ReviewInvalid(field, f"'{field}' is not an accepted field.")
    review = {"table": dict(current["table"]), "columns": [dict(c) for c in current["columns"]]}

    table_edit = edit.get("table", {})
    if not isinstance(table_edit, dict):
        raise ReviewInvalid("table", "'table' must be an object.")
    for key in set(table_edit) - {"description", "null_tokens"}:
        raise ReviewInvalid(f"table.{key}", f"'table.{key}' is not an editable field.")
    if "description" in table_edit:
        review["table"]["description"] = _description(table_edit["description"], "table.description")
    if "null_tokens" in table_edit:
        tokens = table_edit["null_tokens"]
        if (not isinstance(tokens, list) or len(tokens) > MAX_NULL_TOKENS
                or not all(isinstance(t, str) and len(t) <= MAX_NULL_TOKEN_LENGTH for t in tokens)):
            raise ReviewInvalid("table.null_tokens",
                                f"'table.null_tokens' must be a list of at most {MAX_NULL_TOKENS} short strings.")
        review["table"]["null_tokens"] = sorted({t.strip().lower() for t in tokens})

    column_edits = edit.get("columns", [])
    if not isinstance(column_edits, list):
        raise ReviewInvalid("columns", "'columns' must be a list.")
    by_index = {c["index"]: c for c in review["columns"]}
    for position, column_edit in enumerate(column_edits):
        prefix = f"columns[{position}]"
        if not isinstance(column_edit, dict) or not isinstance(column_edit.get("index"), int) \
                or column_edit["index"] not in by_index:
            raise ReviewInvalid(f"{prefix}.index", f"'{prefix}.index' must name a profiled column.")
        column = by_index[column_edit["index"]]
        for key in set(column_edit) - {"index", "identifier", "type", "date_format", "excluded",
                                       "description", "value_hints"}:
            raise ReviewInvalid(f"{prefix}.{key}", f"'{prefix}.{key}' is not an editable field.")
        if "identifier" in column_edit:
            if not is_valid_identifier(column_edit["identifier"]):
                raise ReviewInvalid(f"{prefix}.identifier",
                                    f"'{prefix}.identifier' must be lowercase letters, digits and "
                                    "underscores, not start with a digit, and not be a reserved word.")
            column["identifier"] = column_edit["identifier"]
        if "type" in column_edit:
            if column_edit["type"] not in ALLOWED_TYPES:
                raise ReviewInvalid(f"{prefix}.type", f"'{prefix}.type' must be one of {', '.join(ALLOWED_TYPES)}.")
            column["type"] = column_edit["type"]
            if column["type"] != TYPE_DATE:
                column["date_format"] = None
        if "date_format" in column_edit:
            fmt = column_edit["date_format"]
            if fmt is not None and fmt not in DATE_FORMATS:
                raise ReviewInvalid(f"{prefix}.date_format",
                                    f"'{prefix}.date_format' must be one of {', '.join(DATE_FORMATS)}.")
            column["date_format"] = fmt
        if "excluded" in column_edit:
            if not isinstance(column_edit["excluded"], bool):
                raise ReviewInvalid(f"{prefix}.excluded", f"'{prefix}.excluded' must be true or false.")
            column["excluded"] = column_edit["excluded"]
        if "description" in column_edit:
            column["description"] = _description(column_edit["description"], f"{prefix}.description")
        if "value_hints" in column_edit:
            hints = column_edit["value_hints"]
            known = {h["value"] for h in column["value_hints"]}
            if not isinstance(hints, list) or not all(
                    isinstance(h, dict) and h.get("value") in known and isinstance(h.get("approved"), bool)
                    for h in hints):
                raise ReviewInvalid(f"{prefix}.value_hints",
                                    f"'{prefix}.value_hints' may only approve or drop profiled values.")
            approved = {h["value"]: h["approved"] for h in hints}
            column["value_hints"] = [
                {"value": h["value"], "approved": approved.get(h["value"], h["approved"])}
                for h in column["value_hints"]
            ]
    return review


def _description(value, field: str) -> str:
    if not isinstance(value, str) or len(value) > MAX_DESCRIPTION_LENGTH:
        raise ReviewInvalid(field, f"'{field}' must be text of at most {MAX_DESCRIPTION_LENGTH} characters.")
    return value.strip()


def publish_blockers(review: dict) -> list[dict]:
    """Every unmet publish-gate condition, as `{code, column?}`."""
    blockers: list[dict] = []
    included = [c for c in review["columns"] if not c.get("excluded")]
    if not included:
        blockers.append({"code": NO_INCLUDED_COLUMNS})
    for column in included:
        if column["type"] == TYPE_DATE and column.get("date_format") not in DATE_FORMATS:
            blockers.append({"code": DATE_FORMAT_REQUIRED, "column": column["identifier"]})
    if not (review["table"].get("description") or "").strip():
        blockers.append({"code": DESCRIPTION_REQUIRED})
    seen: set[str] = set()
    for column in included:
        if column["identifier"] in seen:
            blockers.append({"code": DUPLICATE_IDENTIFIER, "column": column["identifier"]})
        seen.add(column["identifier"])
    return blockers


def _castable(review: dict) -> list[dict]:
    """Included columns whose cast is computable (date columns without a
    format are reported by the gate, not the load report)."""
    return [c for c in review["columns"] if not c.get("excluded")
            and not (c["type"] == TYPE_DATE and c.get("date_format") not in DATE_FORMATS)]


def compute_load_report(con, review: dict) -> dict:
    """Rows read, rows that would load, rows that would be rejected (first 100
    with reason `cast_failed`), and a typed 20-row preview of included columns."""
    tokens = [t.lower() for t in review["table"].get("null_tokens", DEFAULT_NULL_TOKENS)]
    params = {"tokens": tokens}
    columns = _castable(review)
    rows_read = con.execute("SELECT COUNT(*) FROM staging").fetchone()[0]
    fail_terms = []
    for c in columns:
        v = _nv(_col(c["index"]))
        if c["type"] != TYPE_TEXT:
            fail_terms.append((c["identifier"], f"({v} IS NOT NULL AND {cast_expr(c['type'], c.get('date_format'), v)} IS NULL)"))
    any_fail = " OR ".join(term for _, term in fail_terms) or "FALSE"
    rejected = _run(con, f"SELECT COUNT(*) FROM staging WHERE {any_fail}", params).fetchone()[0]
    rejects = []
    if rejected:
        first_col = " ".join(f"WHEN {term} THEN '{ident}'" for ident, term in fail_terms)
        rejects = [
            {"row": int(row), "reason": "cast_failed", "column": column}
            for row, column in _run(
                con,
                f"SELECT __row, CASE {first_col} END FROM staging WHERE {any_fail} "
                f"ORDER BY __row LIMIT {MAX_REJECT_ROWS}",
                params,
            ).fetchall()
        ]
    preview = []
    if columns:
        select = ", ".join(
            f"{cast_expr(c['type'], c.get('date_format'), _nv(_col(c['index'])))} AS {c['identifier']}"
            for c in columns
        )
        cursor = _run(
            con,
            f"SELECT {select} FROM staging WHERE NOT ({any_fail}) ORDER BY __row LIMIT {PREVIEW_ROWS}",
            params,
        )
        names = [d[0] for d in cursor.description]
        preview = [{n: _jsonable(val) for n, val in zip(names, row)} for row in cursor.fetchall()]
    return {
        "rows_read": int(rows_read), "rows_to_load": int(rows_read - rejected),
        "rows_rejected": int(rejected), "rejects": rejects, "preview": preview,
    }


def write_parquet(con, review: dict, out_path: str) -> None:
    """Converts the staged table to Parquet with the approved types, formats and
    null tokens, dropping rows that fail any cast."""
    tokens = [t.lower() for t in review["table"].get("null_tokens", DEFAULT_NULL_TOKENS)]
    columns = [c for c in review["columns"] if not c.get("excluded")]
    select_parts, fail_terms = [], []
    for c in columns:
        v = _nv(_col(c["index"]))
        expr = cast_expr(c["type"], c.get("date_format"), v)
        select_parts.append(f"CAST({expr} AS {_duck_type(c['type'])}) AS {c['identifier']}")
        if c["type"] != TYPE_TEXT:
            fail_terms.append(f"({v} IS NOT NULL AND {expr} IS NULL)")
    any_fail = " OR ".join(fail_terms) or "FALSE"
    path = out_path.replace("\\", "/").replace("'", "''")
    _run(
        con,
        f"COPY (SELECT {', '.join(select_parts)} FROM staging WHERE NOT ({any_fail}) ORDER BY __row) "
        f"TO '{path}' (FORMAT PARQUET)",
        {"tokens": tokens},
    )


def build_contract(review: dict, *, display_name: str, sheet: str | None, version: int) -> dict:
    """The canonical published contract: only included columns, with types,
    descriptions and approved value hints."""
    columns = []
    for c in review["columns"]:
        if c.get("excluded"):
            continue
        columns.append({
            "name": c["identifier"], "label": c["label"], "type": c["type"],
            "date_format": c.get("date_format"), "description": c.get("description") or "",
            "value_hints": [h["value"] for h in c.get("value_hints", []) if h.get("approved")],
        })
    return {
        "relation": review["table"]["relation"],
        "description": review["table"]["description"],
        "source_file": display_name, "sheet": sheet, "version": version,
        "columns": columns,
    }


def schema_diff(served_contract: dict | None, review: dict) -> dict | None:
    """Added, removed, retyped and renamed columns of a pending version against
    the served one. Renamed means the same original header under a new
    identifier; it is not also reported as added and removed."""
    if not served_contract:
        return None
    old = {c["name"]: c for c in served_contract.get("columns", [])}
    new = {c["identifier"]: c for c in review["columns"] if not c.get("excluded")}
    old_by_label = {c.get("label"): c["name"] for c in old.values() if c.get("label")}
    added, removed, retyped, renamed = [], [], [], []
    renamed_old: set[str] = set()
    for name, column in new.items():
        if name in old:
            if old[name]["type"] != column["type"]:
                retyped.append({"column": name, "from": old[name]["type"], "to": column["type"]})
            continue
        previous = old_by_label.get(column["label"])
        if previous and previous not in new:
            renamed.append({"from": previous, "to": name})
            renamed_old.add(previous)
        else:
            added.append(name)
    for name in old:
        if name not in new and name not in renamed_old:
            removed.append(name)
    return {"added": sorted(added), "removed": sorted(removed), "retyped": retyped, "renamed": renamed}


def relation_for(display_name: str, sheet: str | None, taken: set[str]) -> str:
    """Relation name from the sheet name (XLSX) or the file stem, de-duplicated
    against the tenant's other files' relations."""
    stem = sheet if sheet else os.path.splitext(display_name)[0]
    base = sanitize_identifier(stem)
    candidate, n = base, 2
    while candidate in taken:
        suffix = f"_{n}"
        candidate = base[: 63 - len(suffix)] + suffix
        n += 1
    return candidate
